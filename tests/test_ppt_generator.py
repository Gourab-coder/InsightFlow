from pathlib import Path

from pptx import Presentation

from ppt_generator import generate_powerpoint


def test_generate_powerpoint_contains_required_slides():
    package = {
        "metadata": {"analysis_type": "Sales"},
        "kpis": {"numeric_columns": {"Revenue": {"sum": 300, "average": 150}}},
        "trends": {},
        "segments": {},
    }
    insights = {
        "executive_summary": "Revenue is stable.",
        "key_findings": [{"finding": "Revenue totals 300"}],
        "recommendations": ["Track monthly revenue"],
    }
    path = Path("outputs/presentations/test_presentation.pptx")
    try:
        generate_powerpoint(package, insights, path)
        presentation = Presentation(path)
        titles = [slide.shapes.title.text for slide in presentation.slides]

        assert path.exists()
        assert len(presentation.slides) == 5
        assert titles[0] == "InsightFlow Business Report"
        assert "Insights and Recommended Actions" in titles
    finally:
        if path.exists():
            path.unlink()
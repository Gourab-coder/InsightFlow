"""PowerPoint generation from verified InsightFlow evidence."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches


def create_presentation() -> Presentation:
	"""Create a presentation with a standard widescreen layout."""
	presentation = Presentation()
	presentation.slide_width = Inches(13.333)
	presentation.slide_height = Inches(7.5)
	return presentation


def _slide(presentation: Presentation, title: str, body: list[str] | None = None):
	slide = presentation.slides.add_slide(presentation.slide_layouts[1])
	slide.shapes.title.text = title
	if body:
		frame = slide.placeholders[1].text_frame
		frame.clear()
		for index, line in enumerate(body):
			paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
			paragraph.text = line
			paragraph.level = 0
	return slide


def add_title_slide(presentation: Presentation, analysis_package: dict[str, Any]) -> None:
	slide = presentation.slides.add_slide(presentation.slide_layouts[0])
	slide.shapes.title.text = "InsightFlow Business Report"
	slide.placeholders[1].text = analysis_package["metadata"].get("analysis_type", "General Business")


def add_summary_slide(presentation: Presentation, ai_output: dict[str, Any]) -> None:
	_slide(presentation, "Executive Summary", [ai_output["executive_summary"]])


def add_kpi_slide(presentation: Presentation, analysis_package: dict[str, Any]) -> None:
	lines = []
	numeric_columns = analysis_package["kpis"].get("numeric_columns", {})
	primary_metrics = analysis_package.get("business_insights", {}).get("primary_metrics", [])
	if primary_metrics:
		numeric_columns = {name: numeric_columns[name] for name in primary_metrics if name in numeric_columns}
	for metric, values in list(numeric_columns.items())[:5]:
		lines.append(f"{metric}: total {values.get('sum')}, average {values.get('average')}")
	trends = analysis_package.get("trends", {})
	if trends.get("periods"):
		latest = trends["periods"][-1]
		lines.append(f"Latest trend period {latest['period']}: {latest['sum']}")
	_slide(presentation, "KPI Overview", lines or ["No numeric KPI columns detected."])


def add_segment_slide(presentation: Presentation, analysis_package: dict[str, Any]) -> None:
	segments = analysis_package.get("segments", {})
	lines = [f"{group['name']}: {group['sum']}" for group in segments.get("groups", [])[:5]]
	trends = analysis_package.get("trends", {})
	if trends.get("periods"):
		lines.append("Trend: " + ", ".join(f"{item['period']}={item['sum']}" for item in trends["periods"][-4:]))
	_slide(presentation, "Segments and Trend", lines or ["No segment or date trend available."])


def add_insights_slide(presentation: Presentation, ai_output: dict[str, Any]) -> None:
	lines = []
	for finding in ai_output["key_findings"][:5]:
		lines.append(str(finding.get("finding", finding) if isinstance(finding, dict) else finding))
	_slide(presentation, "Key Business Insights", lines or ["No additional findings."])


def add_recommendations_slide(presentation: Presentation, ai_output: dict[str, Any]) -> None:
	lines = ["Insights:"]
	for finding in ai_output["key_findings"][:3]:
		lines.append(str(finding.get("finding", finding) if isinstance(finding, dict) else finding))
	lines.append("Recommended actions:")
	lines.extend(str(item) for item in ai_output["recommendations"][:4])
	_slide(presentation, "Insights and Recommended Actions", lines)


def generate_powerpoint(
	analysis_package: dict[str, Any], ai_output: dict[str, Any], output_path: str | Path
) -> Path:
	"""Generate and save the executive PowerPoint."""
	required = ("executive_summary", "key_findings", "recommendations")
	missing = [field for field in required if field not in ai_output]
	if missing:
		raise ValueError(f"Cannot generate presentation; AI output is missing: {', '.join(missing)}")

	presentation = create_presentation()
	add_title_slide(presentation, analysis_package)
	add_summary_slide(presentation, ai_output)
	add_kpi_slide(presentation, analysis_package)
	add_segment_slide(presentation, analysis_package)
	add_recommendations_slide(presentation, ai_output)

	destination = Path(output_path)
	destination.parent.mkdir(parents=True, exist_ok=True)
	presentation.save(destination)
	return destination

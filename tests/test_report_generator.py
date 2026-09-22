from pathlib import Path

from docx import Document

from report_generator import generate_word_report


def test_report_rejects_incomplete_ai_output():
    try:
        generate_word_report({}, {}, Path("outputs/reports/invalid.docx"))
    except ValueError as exc:
        assert "AI output is missing" in str(exc)
    else:
        raise AssertionError("Incomplete AI output should be rejected")


def test_generate_word_report_contains_required_sections():
    package = {
        "metadata": {"rows": 2, "columns": 2, "analysis_type": "Sales", "business_objective": "Review sales"},
        "data_quality": {"missing_values": 0, "duplicate_rows": 0, "warnings": []},
        "kpis": {"numeric_columns": {"Revenue": {"sum": 300, "average": 150, "minimum": 100, "maximum": 200}}},
        "trends": {},
        "segments": {},
    }
    insights = {
        "executive_summary": "Revenue is stable.",
        "key_findings": [{"finding": "Revenue totals 300", "evidence": "KPI summary"}],
        "risks": ["Small sample"],
        "opportunities": ["Review growth"],
        "recommendations": ["Track monthly revenue"],
    }

    path = Path("outputs/reports/test_report.docx")
    try:
        generate_word_report(package, insights, path)
        document = Document(path)
        headings = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading")]

        assert path.exists()
        assert "Executive Summary" in headings
        assert "KPI Summary" in headings
        assert "Recommendations" in headings
        assert any("Revenue" in cell.text for table in document.tables for row in table.rows for cell in row.cells)
    finally:
        if path.exists():
            path.unlink()